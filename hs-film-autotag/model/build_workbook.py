"""Build the pricing and unit-economics workbook (formulas, not pasted results).

Reads measured numbers from results/benchmark_results.json and the adoption simulation output
from results/adoption_grid.csv and results/adoption_trajectory.json, then writes
model/unit_economics.xlsx. Recalculate with LibreOffice before reading cached values.

Usage: python -I build_workbook.py PROJECT_DIR
"""

import csv
import json
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")

USD = '$#,##0;($#,##0);"-"'
USD2 = '$#,##0.00;($#,##0.00);"-"'
USD4 = '$#,##0.0000;($#,##0.0000);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM3 = '#,##0.000;(#,##0.000);"-"'
MULT = '0.0"x"'

R = {}  # key -> absolute cell reference, e.g. "Inputs!$B$5"


def setw(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, cell, value, font=BLACK, fmt=None, fill=None, bold=False, wrap=False):
    c = ws[cell]
    c.value = value
    c.font = Font(name=FONT, color=font.color, bold=bold or font.bold, size=font.size)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


class Block:
    """Label | value | unit | note rows, recording each value cell under a key."""

    def __init__(self, ws, start_row=3):
        self.ws = ws
        self.r = start_row

    def section(self, title):
        self.r += 1
        for col in "ABCD":
            self.ws[f"{col}{self.r}"].fill = SECTION_FILL
        put(self.ws, f"A{self.r}", title, bold=True)
        self.r += 1

    def row(self, key, label, value, unit="", note="", fmt=None, kind="calc", key_assumption=False):
        ws, r = self.ws, self.r
        put(ws, f"A{r}", label)
        font = {"input": BLUE, "calc": BLACK, "link": GREEN}[kind]
        if isinstance(value, str) and value.startswith("="):
            value = value.format(**R)
        put(ws, f"B{r}", value, font=font, fmt=fmt, fill=KEY_FILL if key_assumption else None)
        put(ws, f"C{r}", unit)
        put(ws, f"D{r}", note, wrap=False)
        R[key] = f"'{ws.title}'!$B${r}" if " " in ws.title else f"{ws.title}!$B${r}"
        self.r += 1
        return r


def header_row(ws, row, labels):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font = HEAD
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")


def build(project):
    bench = json.load(open(os.path.join(project, "results", "benchmark_results.json")))
    grid = list(csv.DictReader(open(os.path.join(project, "results", "adoption_grid.csv"))))
    traj = json.load(open(os.path.join(project, "results", "adoption_trajectory.json")))

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    ws_in = wb.create_sheet("Inputs")
    ws_b = wb.create_sheet("Benchmark")
    ws_u = wb.create_sheet("UnitEcon")
    ws_s = wb.create_sheet("Sensitivity")
    ws_m = wb.create_sheet("Monthly")
    ws_mk = wb.create_sheet("Market")
    ws_a = wb.create_sheet("Adoption")

    # ------------------------------------------------------------------ Benchmark (measured)
    setw(ws_b, [44, 16, 16, 16, 16, 18, 60])
    put(ws_b, "A1", "Measured compute benchmark (CPU, no GPU available in the test box)", font=TITLE)
    put(ws_b, "A2", f"Footage: {bench['footage']['file']}, {bench['footage']['license']}, "
                    f"{bench['footage']['resolution']} @ {bench['footage']['fps']} fps. "
                    f"Source: {bench['footage']['source']}")
    b = Block(ws_b, 3)
    b.section("Raw measurements (hardcoded = measured; see hs-film-autotag/bench/bench_pipeline.py)")
    d1 = bench["decode_full_1thread"]
    dn = bench["decode_full_allthreads"]
    b.row("seg_len", "Segment length", round(bench["footage"]["segment_s"], 2), "s of video",
          "ffprobe duration of the downloaded segment", NUM1, "input")
    b.row("dec1_cpu", "Full 1080p H.264 decode, 1 thread: CPU time", round(d1["user_s"] + d1["sys_s"], 3),
          "CPU-s", "ffmpeg -benchmark utime+stime", NUM1, "input")
    b.row("decN_cpu", "Full decode, all 4 threads: CPU time", round(dn["user_s"] + dn["sys_s"], 3),
          "CPU-s", "ffmpeg -benchmark utime+stime", NUM1, "input")
    b.row("decN_wall", "Full decode, all 4 threads: wall time", dn["wall_s"], "s", "ffmpeg -benchmark rtime",
          NUM1, "input")
    seg = bench["segmentation_lowres_5fps"]
    b.row("seg_proc", "Snap-signal processing (5 fps, 320x180, motion-compensated diff)", seg["proc_s"], "CPU-s",
          "Python/OpenCV time excluding decode", NUM3, "input")
    b.row("seg_frames", "Frames processed for snap signal", seg["frames"], "frames", "", NUM, "input")
    b.row("lines_ms", "Field-line detection per 1080p frame", round(bench["field_lines_ms_per_1080p_frame"], 2),
          "ms", "HSV mask + Canny + HoughLinesP, mean of 40 frames", NUM1, "input")
    sh = bench["snap_heuristic"]
    b.row("snaps_true", "Snaps in segment (manual count)", sh["true_snaps_in_segment_manual"], "plays",
          "Counted from 2-second contact sheets", NUM, "input")
    b.row("snaps_tp", "Naive motion heuristic: true snaps found", sh["true_positives"], "plays",
          f"{sh['candidates']} candidates returned", NUM, "input")
    b.row("snap_recall", "Naive snap-heuristic recall", "={snaps_tp}/{snaps_true}", "", "Learned detector needed",
          PCT)
    mr = bench["manual_frame_review"]
    b.row("frames_reviewed", "Pre-snap frames reviewed by eye", mr["presnap_frames"], "frames", "", NUM, "input")
    b.row("frames_numbers", "...with a legible yard number", mr["yard_number_legible"], "frames",
          "Others need camera-motion or game-state continuity", NUM, "input")
    b.row("frames_all22", "...with every player inside the frame", mr["all_players_in_frame"], "frames", "", NUM,
          "input")
    b.section("Derived rates")
    b.row("dec_rate", "Decode cost", "={dec1_cpu}/{seg_len}", "CPU-s per video-s", "1080p, 4 Mbps", NUM3)
    b.row("dec_speed", "Decode speed, 4 vCPU", "={seg_len}/{decN_wall}", "x real time", "", MULT)
    b.row("seg_rate", "Snap-signal cost", "={seg_proc}/{seg_len}", "CPU-s per video-s", "", NUM3)
    b.row("reg_s", "Registration cost per frame", "={lines_ms}/1000", "CPU-s per frame", "", NUM3)

    b.r += 1
    det_hdr = b.r
    header_row(ws_b, det_hdr, ["Detector (key)", "CPU ms/frame, 1 thread", "CPU ms/frame, 4 threads",
                               "Persons found, mean of 8 pre-snap frames", "Published T4 GPU ms/frame",
                               "Tiles", "License and note"])
    det = bench["detector_ms_per_frame"]
    counts = bench["person_counts_on_8_presnap_frames"]
    models = [
        ("yolo11n_640", "yolo11n_640", None, 1.5, 1, "AGPL-3.0 (commercial use needs Ultralytics license). GPU: docs.ultralytics.com/models/yolo11"),
        ("yolo11s_640", "yolo11s_640", "yolo11s_full", 2.5, 1, "AGPL-3.0. GPU: Ultralytics T4 TensorRT10"),
        ("yolo11s_640_tiled3", "yolo11s_640_tiled3", "yolo11s_tiled3", 2.5, 3, "AGPL-3.0. 3 horizontal tiles, 15% overlap"),
        ("yolo11m_640", "yolo11m_640", None, 4.7, 1, "AGPL-3.0"),
        ("rfdetr_nano_384", "rfdetr_nano_384", None, 2.3, 1, "Apache-2.0. GPU: rfdetr.roboflow.com (T4 TensorRT FP16)"),
        ("rfdetr_medium_576", "rfdetr_medium_576", "rfdetr_m_full", 4.4, 1, "Apache-2.0"),
        ("rfdetr_medium_576_tiled3", "rfdetr_medium_576_tiled3", "rfdetr_m_tiled3", 4.4, 3, "Apache-2.0. 3 tiles"),
    ]
    r = det_hdr + 1
    first_det = r
    for key, bkey, ckey, gpu_ms, tiles, note in models:
        put(ws_b, f"A{r}", key)
        put(ws_b, f"B{r}", round(det[f"{bkey}_t1"]["ms_mean"], 1), font=BLUE, fmt=NUM1)
        put(ws_b, f"C{r}", round(det[f"{bkey}_t4"]["ms_mean"], 1), font=BLUE, fmt=NUM1)
        if ckey:
            put(ws_b, f"D{r}", round(sum(counts[ckey]) / len(counts[ckey]), 1), font=BLUE, fmt=NUM1)
        put(ws_b, f"E{r}", gpu_ms, font=BLUE, fmt=NUM1)
        put(ws_b, f"F{r}", tiles, font=BLUE, fmt=NUM)
        put(ws_b, f"G{r}", note)
        r += 1
    last_det = r - 1
    R["det_names"] = f"Benchmark!$A${first_det}:$A${last_det}"
    R["det_ms1"] = f"Benchmark!$B${first_det}:$B${last_det}"
    R["det_gpu"] = f"Benchmark!$E${first_det}:$E${last_det}"
    R["det_tiles"] = f"Benchmark!$F${first_det}:$F${last_det}"
    put(ws_b, f"A{r + 1}", "Person counts include sideline players, coaches and officials; on-field filtering "
                           "needs the field mask. Counts are a sanity check, not an accuracy measure.")
    put(ws_b, f"A{r + 2}", "GPU column is vendor-published latency (not measured here) and feeds only the "
                           "'GPU alternative' line in UnitEcon.")

    # ------------------------------------------------------------------ Inputs
    setw(ws_in, [52, 14, 22, 90])
    put(ws_in, "A1", "Inputs and assumptions (blue = input, yellow = key assumption, green = link)", font=TITLE)
    put(ws_in, "A2", "Change blue cells only. Every other sheet recalculates from here and from Benchmark.")
    s = Block(ws_in, 3)
    s.section("Pricing")
    s.row("price", "Price per program per season", 900, "$/program/season",
          "Model choice. Anchors: MaxPreps Advantage $600/team/season; Balltime HS $499/season; Hudl Silver+ "
          "with Assist $1,500/program/yr (hudl.com/pricing/high-school)", USD, "input", True)
    s.row("pay_fee", "Payment processing and collections", 0.025, "% of price",
          "Assumption: blend of card, ACH and paper checks", PCT, "input")
    s.section("Usage per program per season")
    s.row("own_games", "Own games broken down", 10.5, "games", "Assumption: UIL 10-game regular season plus average playoff run", NUM1, "input")
    s.row("scout_games", "Opponent (scout) games broken down", 15, "games",
          "Assumption: ~1.5 opponent games/week; Hudl says coaches break down 2+ games/week in total", NUM1, "input")
    s.row("other_games", "Scrimmages, spring game", 3, "games", "Assumption", NUM1, "input")
    s.row("plays", "Plays per game (both teams, incl. special teams)", 140, "plays",
          "Benchmark segment: 10 plays in 10 min; one program logged 615 offensive plays in 10 games "
          "(theteamroom.substack.com)", NUM, "input", True)
    s.row("clip_s", "Seconds of video per play clip", 12, "s", "Assumption: pre-snap to whistle", NUM, "input")
    s.row("angles", "Camera angles processed per play", 1, "angles",
          "Sideline only. Set 2 to also process end zone", NUM1, "input")
    s.row("share_cont", "Share of games arriving as continuous (unclipped) film", 0.3, "%",
          "Assumption: tripod/Focus-style recordings vs Hudl cut-ups", PCT, "input")
    s.row("cont_h", "Length of continuous film per game", 2.2, "hours",
          "Benchmark source game ran 8,029 s (2.2 h) including halftime", NUM1, "input")
    s.section("Compute (per play)")
    s.row("det_model", "Detector used", "yolo11s_640_tiled3", "key", "Pick a key from the Benchmark table",
          None, "input", True)
    s.row("det_frames", "Detector frames per play", 9, "frames",
          "3 pre-snap (formation, personnel) + 6 post-snap at 2 fps (play type, direction, result)", NUM, "input")
    s.row("reg_frames", "Registration frames per play", 6, "frames", "Field lines at snap and at play end", NUM, "input")
    s.row("overhead", "Orchestration, retries and I/O overhead", 1.5, "x", "Assumption", MULT, "input")
    s.row("vcpu_price", "vCPU-hour price", "=0.1785/4", "$/vCPU-h",
          "AWS c7i.xlarge on-demand $0.1785/h for 4 vCPU, us-east-1 (calculator.holori.com, Jul 2026)", USD4, "input")
    s.row("gpu_price", "GPU-hour price (T4, for the GPU alternative)", 0.526, "$/GPU-h",
          "AWS g4dn.xlarge on-demand, us-east-1 (calculator.holori.com); spot ~$0.23-0.26", USD4, "input")
    s.row("par_vcpu", "vCPUs per game job (fan-out)", 32, "vCPU",
          "Plays are independent, so a game splits across workers", NUM, "input")
    s.row("uplink", "Coach upload bandwidth", 20, "Mbps", "Assumption; the real latency driver", NUM, "input")
    s.row("bitrate", "Video bitrate", 4.01, "Mbps", "Measured on the benchmark file (ffprobe)", NUM1, "input")
    s.section("Storage and transfer")
    s.row("retention", "Raw video retention after processing", 14, "days", "Design rule: delete raw video", NUM, "input")
    s.row("stor_price", "Object storage price", 0.023, "$/GB-month", "S3 Standard list price (assumption)", USD4, "input")
    s.row("egress_gb", "Video served back to coaches per game", 0, "GB",
          "0 if tags are written back to Hudl and film is watched there", NUM1, "input")
    s.row("egress_price", "Egress price", 0.09, "$/GB", "Assumption: S3 internet egress first tier", USD2, "input")
    s.section("Support and quality assurance")
    s.row("support_h", "Support hours per program per season", 4, "hours", "Assumption", NUM1, "input", True)
    s.row("support_rate", "Support cost per hour", 40, "$/hour", "Assumption: contractor rate", USD, "input")
    s.row("onboard_h", "Onboarding hours for a new program", 3, "hours", "Terminology mapping, first upload", NUM1, "input")
    s.row("qa_min", "Human QA minutes per game", 0, "min/game",
          "Set 10-20 to model a human-in-the-loop guarantee", NUM, "input", True)
    s.row("qa_rate", "QA cost per hour", 20, "$/hour", "Assumption", USD, "input")
    s.section("Acquisition and retention")
    s.row("cac", "Customer acquisition cost per new program", 400, "$",
          "Assumption: founder-led; booth, travel and trial costs. Founder time not counted", USD, "input", True)
    s.row("churn", "Annual churn", 0.2, "%/season", "Assumption; adoption sim draws 8-25% plus value-driven exits", PCT,
          "input", True)
    s.row("disc", "Discount rate", 0.1, "%/yr", "Assumption", PCT, "input")
    s.section("Fixed costs")
    s.row("fixed_m", "Fixed monthly costs (hosting floor, tools, domains)", 250, "$/month", "Assumption", USD, "input")
    s.row("fixed_y", "Annual fixed costs (insurance, coaching-school attendance)", 3500, "$/year, paid in July",
          "Assumption", USD, "input")
    s.row("legal", "One-time legal (DPA addendum, terms, entity)", 3000, "$ in Jan 2027", "Assumption", USD, "input")
    s.section("Season sizes (paying programs; 2027 = free pilots)")
    s.row("pilots_2027", "Free pilot programs, fall 2027", 8, "programs", "Founder capacity, not demand", NUM, "input")
    s.row("scenario", "Adoption scenario (1 = P10, 2 = P50, 3 = P90)", 2, "", "Feeds Monthly from Adoption sheet",
          NUM, "input", True)
    s.section("Seasonality: share of a season's games processed in each month")
    weights = [0, 0, 0.02, 0.05, 0.03, 0.01, 0.01, 0.10, 0.25, 0.27, 0.20, 0.06]
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    w_first = s.r
    for mname, w in zip(months, weights):
        s.row(f"w_{mname}", f"{mname}", w, "% of season",
              "Spring ball Apr-May (5A/6A), season late Aug-Nov, playoffs Nov-Dec" if mname == "Apr" else "",
              PCT, "input")
    w_last = s.r - 1
    R["weights"] = f"Inputs!$B${w_first}:$B${w_last}"
    s.row("w_sum", "Check: weights sum to 100%", f"=SUM(B{w_first}:B{w_last})", "", "", PCT)
    dv = DataValidation(type="list", formula1=f"=Benchmark!$A${first_det}:$A${last_det}", allow_blank=False)
    ws_in.add_data_validation(dv)
    dv.add(R["det_model"].split("!")[1].replace("$", ""))

    # ------------------------------------------------------------------ UnitEcon
    setw(ws_u, [56, 16, 22, 70])
    put(ws_u, "A1", "Unit economics per game and per program-season", font=TITLE)
    u = Block(ws_u, 2)
    u.section("Per game: compute and media")
    u.row("games", "Games per program per season", "={own_games}+{scout_games}+{other_games}", "games", "", NUM1)
    u.row("video_s", "Video seconds processed per game",
          "=(1-{share_cont})*{plays}*{clip_s}*{angles}+{share_cont}*{cont_h}*3600", "s", "", NUM)
    u.row("dec_rate_l", "Decode cost (from Benchmark)", "={dec_rate}", "CPU-s per video-s", "", NUM3, "link")
    u.row("seg_rate_l", "Snap-signal cost (from Benchmark)", "={seg_rate}", "CPU-s per video-s", "", NUM3, "link")
    u.row("det_s", "Detector CPU-s per frame (selected model)",
          "=INDEX({det_ms1},MATCH({det_model},{det_names},0))/1000", "CPU-s per frame", "1-thread measurement", NUM3,
          "link")
    u.row("reg_s_l", "Registration CPU-s per frame", "={reg_s}", "CPU-s per frame", "", NUM3, "link")
    u.row("cpu_dec", "CPU-s: decode", "={video_s}*{dec_rate_l}", "CPU-s", "", NUM)
    u.row("cpu_seg", "CPU-s: play segmentation", "={video_s}*{seg_rate_l}", "CPU-s", "", NUM)
    u.row("cpu_det", "CPU-s: player detection", "={plays}*{det_frames}*{det_s}", "CPU-s", "", NUM)
    u.row("cpu_reg", "CPU-s: field registration", "={plays}*{reg_frames}*{reg_s_l}", "CPU-s", "", NUM)
    u.row("cpu_tot", "CPU-s per game incl. overhead", "=({cpu_dec}+{cpu_seg}+{cpu_det}+{cpu_reg})*{overhead}",
          "CPU-s", "", NUM)
    u.row("cpu_s_play", "CPU-s per play", "={cpu_tot}/{plays}", "CPU-s per play", "", NUM1)
    u.row("compute_game", "Compute cost per game (CPU only)", "={cpu_tot}/3600*{vcpu_price}", "$/game",
          "Target: under $1", USD4, key_assumption=True)
    u.row("target_cost", "Target check: compute under $1 per game", '=IF({compute_game}<1,"PASS","FAIL")', "", "")
    u.row("gpu_det_game", "GPU alternative: detection on T4 (vendor latency)",
          "={plays}*{det_frames}*INDEX({det_gpu},MATCH({det_model},{det_names},0))"
          "*INDEX({det_tiles},MATCH({det_model},{det_names},0))/1000/3600*{gpu_price}",
          "$/game", "Derived, not measured: no GPU in the test box", USD4)
    u.row("storage_game", "Storage per game", "={video_s}*{bitrate}/8/1000*{stor_price}*{retention}/30", "$/game",
          "", USD4)
    u.row("egress_game", "Egress per game", "={egress_gb}*{egress_price}", "$/game", "", USD4)
    u.row("media_game", "Compute + storage + egress per game", "={compute_game}+{storage_game}+{egress_game}",
          "$/game", "", USD4)
    u.section("Latency check (target: breakdown within 10 min of upload)")
    u.row("proc_min", "Processing wall time with fan-out", "={cpu_tot}/{par_vcpu}/60", "minutes", "", NUM1)
    u.row("upload_min", "Upload time for the same video", "={video_s}*{bitrate}/{uplink}/60", "minutes",
          "Happens before the 10-minute clock starts, but coaches feel it", NUM1)
    u.row("target_lat", "Target check: processing under 10 minutes", '=IF({proc_min}<=10,"PASS","FAIL")', "", "")
    u.section("Per program per season")
    u.row("rev", "Revenue", "={price}", "$", "", USD, "link")
    u.row("fees", "Payment processing", "={rev}*{pay_fee}", "$", "", USD)
    u.row("media_season", "Compute, storage, egress", "={media_game}*{games}", "$", "", USD2)
    u.row("support_season", "Support", "={support_h}*{support_rate}", "$", "", USD)
    u.row("qa_season", "Human QA", "={qa_min}/60*{qa_rate}*{games}", "$", "", USD)
    u.row("cogs", "Cost of revenue", "={fees}+{media_season}+{support_season}+{qa_season}", "$", "", USD)
    u.row("gp", "Gross profit", "={rev}-{cogs}", "$", "", USD)
    u.row("gm", "Gross margin", "=IF({rev}>0,{gp}/{rev},0)", "%", "", PCT, key_assumption=True)
    u.row("compute_share", "Compute share of revenue", "=IF({rev}>0,{media_season}/{rev},0)", "%", "", PCT)
    u.section("Acquisition and lifetime value")
    u.row("first_contrib", "First-season contribution after CAC and onboarding",
          "={gp}-{cac}-{onboard_h}*{support_rate}", "$", "", USD)
    u.row("ltv", "Lifetime gross profit (LTV)", "={gp}*(1+{disc})/({disc}+{churn})", "$",
          "Seasonal periods; GP x (1+d)/(d+churn)", USD)
    u.row("ltv_cac", "LTV / CAC", "=IF({cac}>0,{ltv}/{cac},0)", "x", "Rule of thumb: above 3x", MULT)
    u.row("payback", "CAC payback", "=IF({gp}>0,{cac}/{gp},0)", "seasons", "", NUM1)
    u.row("breakeven_prog", "Programs to cover fixed costs", "=({fixed_m}*12+{fixed_y})/{gp}", "programs", "", NUM1)

    # ------------------------------------------------------------------ Sensitivity
    setw(ws_s, [34, 14, 14, 14, 14, 14, 14])
    put(ws_s, "A1", "Sensitivity tables (each cell is a live formula on Inputs/UnitEcon)", font=TITLE)
    # 1: gross margin by price x compute multiplier
    prices = [300, 600, 900, 1200, 1800]
    put(ws_s, "A3", "Gross margin: price (columns) x compute cost multiplier (rows)", bold=True)
    header_row(ws_s, 4, ["Compute multiplier \\ Price"] + [f"${p:,}" for p in prices])
    for j, p in enumerate(prices):
        put(ws_s, f"{get_column_letter(j + 2)}5", p, font=BLUE, fmt=USD)
    put(ws_s, "A5", "price ->", font=BLACK)
    mults = [1, 10, 30, 100, 300]
    for i, m in enumerate(mults):
        rr = 6 + i
        put(ws_s, f"A{rr}", m, font=BLUE, fmt=MULT)
        for j in range(len(prices)):
            col = get_column_letter(j + 2)
            f = (f"=({col}$5*(1-{R['pay_fee']})-({R['compute_game']}*$A{rr}+{R['storage_game']}+{R['egress_game']})"
                 f"*{R['games']}-{R['support_season']}-{R['qa_season']})/{col}$5")
            put(ws_s, f"{col}{rr}", f, fmt=PCT)
    put(ws_s, "A11", "Reading: compute would have to rise ~100x before it dents margin; support is the real cost.")
    # 2: LTV/CAC by churn x CAC
    put(ws_s, "A13", "LTV / CAC: CAC (columns) x annual churn (rows)", bold=True)
    cacs = [200, 400, 800, 1600, 3200]
    header_row(ws_s, 14, ["Churn \\ CAC"] + [f"${c:,}" for c in cacs])
    put(ws_s, "A15", "CAC ->")
    for j, c in enumerate(cacs):
        put(ws_s, f"{get_column_letter(j + 2)}15", c, font=BLUE, fmt=USD)
    churns = [0.1, 0.2, 0.3, 0.4, 0.5]
    for i, ch in enumerate(churns):
        rr = 16 + i
        put(ws_s, f"A{rr}", ch, font=BLUE, fmt=PCT)
        for j in range(len(cacs)):
            col = get_column_letter(j + 2)
            f = f"={R['gp']}*(1+{R['disc']})/({R['disc']}+$A{rr})/{col}$15"
            put(ws_s, f"{col}{rr}", f, fmt=MULT)
    # 3: compute $/game by detector x frames per play
    put(ws_s, "A23", "Compute $ per game: detector (rows) x detector frames per play (columns)", bold=True)
    fr = [3, 9, 30, 60, 300]
    header_row(ws_s, 24, ["Detector \\ frames per play"] + [str(x) for x in fr])
    put(ws_s, "A25", "frames ->")
    for j, x in enumerate(fr):
        put(ws_s, f"{get_column_letter(j + 2)}25", x, font=BLUE, fmt=NUM)
    for i, (key, *_rest) in enumerate(models):
        rr = 26 + i
        put(ws_s, f"A{rr}", key, font=BLUE)
        for j in range(len(fr)):
            col = get_column_letter(j + 2)
            f = (f"=({R['cpu_dec']}+{R['cpu_seg']}+{R['cpu_reg']}+{R['plays']}*{col}$25"
                 f"*INDEX({R['det_ms1']},MATCH($A{rr},{R['det_names']},0))/1000)*{R['overhead']}/3600*{R['vcpu_price']}")
            put(ws_s, f"{col}{rr}", f, fmt=USD2)
    put(ws_s, f"A{26 + len(models)}", "60 frames/play = full tracking at 10 fps for 6 s; 300 = 30 fps for 10 s. "
                                      "CPU only, on-demand prices.")
    # 4: break-even programs by price x annual fixed burden
    put(ws_s, "A36", "Programs needed to cover an annual cost base: price (columns) x annual cost (rows)", bold=True)
    header_row(ws_s, 37, ["Annual cost \\ Price"] + [f"${p:,}" for p in prices])
    put(ws_s, "A38", "price ->")
    for j, p in enumerate(prices):
        put(ws_s, f"{get_column_letter(j + 2)}38", p, font=BLUE, fmt=USD)
    bases = [6500, 25000, 100000, 200000, 400000]
    for i, base in enumerate(bases):
        rr = 39 + i
        put(ws_s, f"A{rr}", base, font=BLUE, fmt=USD)
        for j in range(len(prices)):
            col = get_column_letter(j + 2)
            gp_p = (f"({col}$38*(1-{R['pay_fee']})-{R['media_season']}-{R['support_season']}-{R['qa_season']})")
            put(ws_s, f"{col}{rr}", f"=IF({gp_p}>0,$A{rr}/{gp_p},0)", fmt=NUM)
    put(ws_s, "A44", "Rows: $6.5k = current fixed costs; $100k-$200k = one funded engineer; $400k = small team.")

    # ------------------------------------------------------------------ Adoption (simulation output)
    setw(ws_a, [30, 14, 14, 14, 14, 14, 60])
    put(ws_a, "A1", "Adoption simulation output (hs-film-autotag/sim/adoption_sim.py, seed 20261023, 1,000 draws)",
        font=TITLE)
    put(ws_a, "A2", "Model of stated priors, not evidence. Texas UIL 11-man programs = "
                    f"{traj['total_programs']}. Hardcoded values below are simulation outputs.")
    put(ws_a, "A4", "Paying programs by season, base case ($900, 95% core accuracy)", bold=True)
    header_row(ws_a, 5, ["Season", "P10, Hudl never automates", "P50", "P90",
                         "P50 if Hudl automates 2029", "P50 if Hudl automates 2028"])
    seasons = ["2027", "2028", "2029", "2030"]
    tj = traj["trajectory_p10_p50_p90"]
    for i, sn in enumerate(seasons):
        rr = 6 + i
        put(ws_a, f"A{rr}", sn)
        put(ws_a, f"B{rr}", tj["never"][sn][0], font=BLUE, fmt=NUM)
        put(ws_a, f"C{rr}", tj["never"][sn][1], font=BLUE, fmt=NUM)
        put(ws_a, f"D{rr}", tj["never"][sn][2], font=BLUE, fmt=NUM)
        put(ws_a, f"E{rr}", tj["2029"][sn][1], font=BLUE, fmt=NUM)
        put(ws_a, f"F{rr}", tj["2028"][sn][1], font=BLUE, fmt=NUM)
    put(ws_a, "A10", "2027 counts are free-pilot uptake in the sim; Monthly uses Inputs' founder-capacity pilots instead.")
    R["adopt_tbl"] = "Adoption!$B$6:$D$9"
    put(ws_a, "A12", "2030 paying programs (P50) by price and accuracy, Hudl never automates", bold=True)
    accs = ["0.85", "0.9", "0.93", "0.95", "0.97"]
    header_row(ws_a, 13, ["Core-column accuracy \\ Price"] + [f"${p:,}" for p in prices])
    for i, a in enumerate(accs):
        rr = 14 + i
        put(ws_a, f"A{rr}", float(a), font=BLUE, fmt=PCT)
        for j, p in enumerate(prices):
            hit = [g for g in grid if g["hudl_automates"] == "never" and int(g["price"]) == p
                   and abs(float(g["accuracy"]) - float(a)) < 1e-9][0]
            put(ws_a, f"{get_column_letter(j + 2)}{rr}", int(hit["adopters_2030_p50"]), font=BLUE, fmt=NUM)
    put(ws_a, "A20", "2030 ARR (P50, $) = programs x price", bold=True)
    header_row(ws_a, 21, ["Core-column accuracy \\ Price"] + [f"${p:,}" for p in prices])
    for i in range(len(accs)):
        rr = 22 + i
        put(ws_a, f"A{rr}", f"=A{14 + i}", fmt=PCT)
        for j, p in enumerate(prices):
            col = get_column_letter(j + 2)
            put(ws_a, f"{col}{rr}", f"={col}{14 + i}*{p}", fmt=USD)
    put(ws_a, "A28", "Same grid if Hudl ships automated football tagging in 2028: P50 programs", bold=True)
    header_row(ws_a, 29, ["Core-column accuracy \\ Price"] + [f"${p:,}" for p in prices])
    for i, a in enumerate(accs):
        rr = 30 + i
        put(ws_a, f"A{rr}", float(a), font=BLUE, fmt=PCT)
        for j, p in enumerate(prices):
            hit = [g for g in grid if g["hudl_automates"] == "2028" and int(g["price"]) == p
                   and abs(float(g["accuracy"]) - float(a)) < 1e-9][0]
            put(ws_a, f"{get_column_letter(j + 2)}{rr}", int(hit["adopters_2030_p50"]), font=BLUE, fmt=NUM)

    # ------------------------------------------------------------------ Monthly cash model
    put(ws_m, "A1", "Monthly cash model, Jan 2027 - Dec 2030 (cash basis; annual billing in August)", font=TITLE)
    cols = ["Month", "Year", "Month #", "Seasonality weight", "Paying programs (season)", "Pilot programs",
            "New paying programs", "Revenue (cash)", "Media COGS", "Support + QA", "Payment fees", "CAC",
            "Fixed costs", "Net cash flow", "Cumulative cash"]
    header_row(ws_m, 3, cols)
    setw(ws_m, [12, 8, 9, 13, 14, 11, 13, 13, 12, 12, 12, 11, 11, 13, 14])
    put(ws_m, "Q3", "Season lookup", bold=True)
    for i, yr in enumerate([2027, 2028, 2029, 2030]):
        put(ws_m, f"Q{4 + i}", yr, font=BLUE)
        if yr == 2027:
            put(ws_m, f"R{4 + i}", 0, font=BLUE, fmt=NUM)
        else:
            put(ws_m, f"R{4 + i}", f"=INDEX({R['adopt_tbl']},{i + 1},{R['scenario']})", font=GREEN, fmt=NUM)
    put(ws_m, "R3", "Paying programs", bold=True)
    row0 = 4
    for k in range(48):
        rr = row0 + k
        y = 2027 + k // 12
        mth = k % 12 + 1
        put(ws_m, f"A{rr}", date(y, mth, 1), fmt="mmm yyyy")
        put(ws_m, f"B{rr}", f"=YEAR(A{rr})", fmt="0")
        put(ws_m, f"C{rr}", f"=MONTH(A{rr})", fmt="0")
        put(ws_m, f"D{rr}", f"=INDEX({R['weights']},C{rr})", fmt=PCT, font=GREEN)
        put(ws_m, f"E{rr}", f"=INDEX($R$4:$R$7,MATCH(B{rr},$Q$4:$Q$7,0))", fmt=NUM)
        put(ws_m, f"F{rr}", f"=IF(B{rr}=2027,{R['pilots_2027']},0)", fmt=NUM)
        prev = (f"IF(B{rr}=2027,0,INDEX($R$4:$R$7,MATCH(B{rr}-1,$Q$4:$Q$7,0)))")
        put(ws_m, f"G{rr}", f"=MAX(0,E{rr}-{prev}*(1-{R['churn']}))", fmt=NUM1)
        put(ws_m, f"H{rr}", f"=IF(C{rr}=8,E{rr}*{R['price']},0)", fmt=USD)
        put(ws_m, f"I{rr}", f"=(E{rr}+F{rr})*{R['media_season']}*D{rr}", fmt=USD2)
        put(ws_m, f"J{rr}", f"=(E{rr}+F{rr})*({R['support_season']}+{R['qa_season']})*D{rr}"
                            f"+IF(C{rr}=8,G{rr}*{R['onboard_h']}*{R['support_rate']},0)", fmt=USD)
        put(ws_m, f"K{rr}", f"=H{rr}*{R['pay_fee']}", fmt=USD)
        put(ws_m, f"L{rr}", f"=IF(AND(C{rr}>=4,C{rr}<=7),G{rr}*{R['cac']}/4,0)", fmt=USD)
        put(ws_m, f"M{rr}", f"={R['fixed_m']}+IF(C{rr}=7,{R['fixed_y']},0)+IF(AND(B{rr}=2027,C{rr}=1),{R['legal']},0)",
            fmt=USD)
        put(ws_m, f"N{rr}", f"=H{rr}-I{rr}-J{rr}-K{rr}-L{rr}-M{rr}", fmt=USD)
        put(ws_m, f"O{rr}", f"=N{rr}" if k == 0 else f"=O{rr - 1}+N{rr}", fmt=USD)
    last = row0 + 47
    put(ws_m, f"A{last + 2}", "Totals", bold=True)
    for col in "HIJKLMN":
        put(ws_m, f"{col}{last + 2}", f"=SUM({col}{row0}:{col}{last})", fmt=USD, bold=True)
    put(ws_m, f"A{last + 3}", "Lowest cumulative cash", bold=True)
    put(ws_m, f"O{last + 3}", f"=MIN(O{row0}:O{last})", fmt=USD, bold=True)
    put(ws_m, f"A{last + 4}", "Founder time is unpaid and not in these numbers (about 1,000 hours a year at 20 h/week).")

    # ------------------------------------------------------------------ Market
    setw(ws_mk, [46, 14, 16, 16, 16, 16, 60])
    put(ws_mk, "A1", "Market sizing (bottom-up, formulas)", font=TITLE)
    header_row(ws_mk, 3, ["Segment", "Programs", "Price/season", "Revenue at 100%", "At 5%", "At 15%", "Source"])
    segs = [
        ("U.S. 11-player HS football (2024-25)", 14269, "=Inputs!$B$5",
         "NFHS 2024-25 participation survey (nfhs.org)"),
        ("Texas UIL 11-man (2024-26)", 1104, "=Inputs!$B$5", "UIL alignment PDF, counted by school"),
        ("Texas UIL six-man (2024-26)", 154, 400, "UIL alignment PDF; six-man needs a separate model"),
        ("NCAA D2 + D3 + NAIA (2025)", 161 + 241 + 94, 2500,
         "Wikipedia 2025 season pages: 161 + 241 + 94; price assumption"),
    ]
    for i, (name, n, p, src) in enumerate(segs):
        rr = 4 + i
        put(ws_mk, f"A{rr}", name)
        put(ws_mk, f"B{rr}", n, font=BLUE, fmt=NUM)
        if isinstance(p, str):
            put(ws_mk, f"C{rr}", R["price"].join(["=", ""]), font=GREEN, fmt=USD)
        else:
            put(ws_mk, f"C{rr}", p, font=BLUE, fmt=USD)
        put(ws_mk, f"D{rr}", f"=B{rr}*C{rr}", fmt=USD)
        put(ws_mk, f"E{rr}", f"=D{rr}*0.05", fmt=USD)
        put(ws_mk, f"F{rr}", f"=D{rr}*0.15", fmt=USD)
        put(ws_mk, f"G{rr}", src)
    put(ws_mk, "A9", "Top-down ceiling: all-sport Hudl spend if every school bought a department package", bold=True)
    put(ws_mk, "A10", "Low package price (Dixie Heights HS, KY, 2022)")
    put(ws_mk, "B10", 8700, font=BLUE, fmt=USD)
    put(ws_mk, "A11", "High package price (Tulsa PS, OK, 2026, per school)")
    put(ws_mk, "B11", "=139500/9", font=BLUE, fmt=USD)
    put(ws_mk, "A12", "Implied all-sport spend, low")
    put(ws_mk, "B12", "=B4*B10", fmt=USD)
    put(ws_mk, "A13", "Implied all-sport spend, high")
    put(ws_mk, "B13", "=B4*B11", fmt=USD)

    # ------------------------------------------------------------------ README
    setw(readme, [110])
    lines = [
        ("Auto-tagging HS football film: pricing and unit economics", TITLE),
        ("Built by hs-film-autotag/model/build_workbook.py on " + date.today().isoformat(), BLACK),
        ("", BLACK),
        ("Colour code: blue = hardcoded input or measurement; black = formula; green = link to another sheet; "
         "yellow fill = key assumption.", BLACK),
        ("Edit only blue cells on Inputs (and Benchmark if you re-run the benchmark). Everything else recalculates.",
         BLACK),
        ("", BLACK),
        ("Sheets", BOLD),
        ("Inputs: prices, usage, compute, storage, support, CAC, churn, fixed costs, seasonality.", BLACK),
        ("Benchmark: measured CPU timings on a CC BY-NC-SA 1080p high school game (no GPU in the test box).", BLACK),
        ("UnitEcon: cost per game, latency check, gross margin, LTV, CAC payback.", BLACK),
        ("Sensitivity: four two-way tables (price x compute, churn x CAC, detector x frames, price x cost base).",
         BLACK),
        ("Monthly: 48-month cash model with football seasonality and August billing.", BLACK),
        ("Market: bottom-up segment sizes.", BLACK),
        ("Adoption: outputs of the Monte Carlo adoption model (priors, not evidence).", BLACK),
        ("", BLACK),
        ("Main caveat: accuracy on core columns is not measured here. It is the largest unknown and is what the "
         "90-day plan tests.", BOLD),
    ]
    for i, (txt, font) in enumerate(lines, 1):
        put(readme, f"A{i}", txt, font=font)

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = True
        ws.freeze_panes = None
    ws_m.freeze_panes = "B4"
    out = os.path.join(project, "model", "unit_economics.xlsx")
    wb.save(out)
    print(out)


if __name__ == "__main__":
    build(sys.argv[1])
